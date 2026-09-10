"""Read the text out of a document a user has uploaded.

The file arrives as bytes, is read in memory and is never written to disk or
stored here: the text goes back to the browser, which sends it with the user's
next question. Nothing about the upload outlives the request.

What can be read, and what cannot - stated to the user rather than papered over:

  PDF with a text layer, Word .docx, and plain text files are read.

  Scanned PDFs and images carry no text layer. Reading them needs OCR, which
  this service does not run, so they come back unreadable with a message that
  says so. Returning an empty string as if it were the document would let the
  assistant answer about a file it never saw.

  Legacy .doc is a binary format python-docx cannot open; the message asks for
  .docx or PDF.

Text is capped at MAX_CHARS. A long test report is still useful from its first
pages, and the answer model has a finite context; the response says when it was
cut so the user knows the assistant did not read the whole thing.
"""

from __future__ import annotations

import base64
import binascii
import io
import logging
import re
from dataclasses import dataclass
from pathlib import PurePath

log = logging.getLogger(__name__)

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 20_000

TEXT_SUFFIXES = {".txt", ".md", ".csv"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}


class ExtractError(ValueError):
    """The upload itself is unusable: not base64, empty, or too large."""


@dataclass
class Extracted:
    filename: str
    kind: str
    text: str
    pages: int | None
    truncated: bool
    readable: bool
    message: str | None = None

    @property
    def characters(self) -> int:
        return len(self.text)


def decode(content_base64: str) -> bytes:
    try:
        data = base64.b64decode(content_base64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ExtractError("the file content is not valid base64") from exc
    if not data:
        raise ExtractError("the file is empty")
    if len(data) > MAX_BYTES:
        raise ExtractError(f"the file is larger than {MAX_BYTES // (1024 * 1024)} MB")
    return data


def _clip(text: str) -> tuple[str, bool]:
    """Tidy whitespace, then cap the length."""
    text = "\n".join(line.rstrip() for line in text.splitlines())
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) > MAX_CHARS:
        return text[:MAX_CHARS], True
    return text, False


def _pdf(data: bytes) -> tuple[str, int]:
    import pymupdf

    with pymupdf.open(stream=data, filetype="pdf") as document:
        pages = document.page_count
        text = "\n\n".join(page.get_text() for page in document)
    return text, pages


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    # Test reports keep their results in tables; skipping them would drop the
    # part of the document the user most wants read.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _unreadable(filename: str, kind: str, message: str, pages: int | None = None) -> Extracted:
    return Extracted(filename, kind, "", pages, False, False, message)


def extract(filename: str, data: bytes) -> Extracted:
    suffix = PurePath(filename).suffix.lower()

    if suffix in IMAGE_SUFFIXES:
        return _unreadable(
            filename,
            "image",
            "Text can't be read from images yet. Upload the PDF or Word version if you have one.",
        )
    if suffix == ".doc":
        return _unreadable(
            filename, "doc", "Older .doc files can't be read. Save it as .docx or PDF and upload that."
        )
    if suffix not in {".pdf", ".docx", *TEXT_SUFFIXES}:
        return _unreadable(
            filename,
            suffix.lstrip(".") or "unknown",
            "This file type can't be read. Upload a PDF, Word (.docx) or text file.",
        )
    if suffix == ".pdf" and not data.startswith(b"%PDF"):
        return _unreadable(filename, "pdf", "This file is named .pdf but is not a PDF.")

    pages: int | None = None
    try:
        if suffix == ".pdf":
            raw, pages = _pdf(data)
            kind = "pdf"
        elif suffix == ".docx":
            raw, kind = _docx(data), "docx"
        else:
            raw, kind = data.decode("utf-8", errors="replace"), "text"
    except Exception as exc:
        log.warning("could not read %s: %s", filename, str(exc)[:160])
        return _unreadable(
            filename,
            suffix.lstrip("."),
            "The file could not be opened. It may be damaged or password-protected.",
        )

    text, truncated = _clip(raw)
    if not text:
        message = (
            "No text was found. It looks like a scanned document, which can't be read yet."
            if kind == "pdf"
            else "No text was found in this file."
        )
        return _unreadable(filename, kind, message, pages)

    return Extracted(filename, kind, text, pages, truncated, True)
