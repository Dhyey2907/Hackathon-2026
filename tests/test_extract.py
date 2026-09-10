"""Tests for reading uploaded documents.

Offline, with the documents built in memory. The cases that matter most are
the ones that must come back *unreadable* - a scanned PDF, an image, a file
misnamed as a PDF - because an empty string passed along as if it were the
document would let the assistant answer about a file it never read.
"""

import base64
import io

import pytest
from fastapi import HTTPException

from bis.api import routes_extract
from bis.documents import extract as ex


def pdf_bytes(text: str | None) -> bytes:
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    if text:
        page.insert_text((72, 72), text)
    data = document.tobytes()
    document.close()
    return data


def docx_bytes(paragraphs: list[str], table_row: list[str] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table_row:
        table = document.add_table(rows=1, cols=len(table_row))
        for cell, value in zip(table.rows[0].cells, table_row, strict=True):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- what can be read -----------------------------------------------------


def test_a_pdf_with_a_text_layer_is_read():
    result = ex.extract("report.pdf", pdf_bytes("Luminous efficacy 110 lm/W per IS 16102"))
    assert result.readable
    assert result.kind == "pdf"
    assert result.pages == 1
    assert "IS 16102" in result.text


def test_a_docx_is_read_including_its_tables():
    """Test reports keep their results in tables."""
    result = ex.extract(
        "report.docx",
        docx_bytes(["Test report for LED lamps"], ["Insulation resistance", "Pass"]),
    )
    assert result.readable
    assert "Test report for LED lamps" in result.text
    assert "Insulation resistance | Pass" in result.text


def test_plain_text_is_read():
    result = ex.extract("notes.txt", b"Licence CM/L-1234567")
    assert result.readable
    assert result.text == "Licence CM/L-1234567"


def test_long_text_is_cut_and_says_so():
    result = ex.extract("long.txt", ("x" * (ex.MAX_CHARS + 500)).encode())
    assert result.truncated
    assert len(result.text) == ex.MAX_CHARS


# --- what must come back unreadable ----------------------------------------


def test_a_scanned_pdf_is_unreadable_not_empty():
    """No text layer. Passing "" along as the document would let the assistant
    answer about a file it never saw."""
    result = ex.extract("scan.pdf", pdf_bytes(None))
    assert not result.readable
    assert "scanned" in (result.message or "")


def test_an_image_is_unreadable_with_a_reason():
    result = ex.extract("licence.jpg", b"\xff\xd8\xff\xe0 not really a jpeg")
    assert not result.readable
    assert result.kind == "image"
    assert result.message


def test_a_file_misnamed_as_pdf_is_caught():
    result = ex.extract("report.pdf", b"<html>not a pdf</html>")
    assert not result.readable
    assert "not a PDF" in (result.message or "")


def test_legacy_doc_asks_for_docx():
    result = ex.extract("report.doc", b"\xd0\xcf\x11\xe0")
    assert not result.readable
    assert ".docx" in (result.message or "")


def test_an_unknown_type_is_unreadable():
    assert not ex.extract("archive.zip", b"PK\x03\x04").readable


def test_a_corrupt_docx_does_not_raise():
    result = ex.extract("broken.docx", b"PK\x03\x04 truncated")
    assert not result.readable
    assert result.message


# --- decoding and the route -----------------------------------------------


def test_bad_base64_is_rejected():
    with pytest.raises(ex.ExtractError):
        ex.decode("not base64 !!!")


def test_an_oversized_file_is_rejected():
    big = base64.b64encode(b"x" * (ex.MAX_BYTES + 1)).decode()
    with pytest.raises(ex.ExtractError):
        ex.decode(big)


def test_the_route_turns_a_decode_error_into_a_400():
    with pytest.raises(HTTPException) as caught:
        routes_extract.post_extract(
            routes_extract.ExtractRequest(filename="a.pdf", content_base64="!!!")
        )
    assert caught.value.status_code == 400


def test_the_route_reads_a_real_file():
    payload = base64.b64encode(pdf_bytes("Scheme-I licence application")).decode()
    response = routes_extract.post_extract(
        routes_extract.ExtractRequest(filename="a.pdf", content_base64=payload)
    )
    assert response.readable
    assert response.characters == len(response.text)
    assert "Scheme-I" in response.text
